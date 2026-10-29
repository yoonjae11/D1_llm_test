"""The D1 planner/query few-shots must be valid inputs for the real pipeline.

The examples are synthetic (invented users, dates and activities). These tests
check that every planner example passes plan validation, every KQL example
passes the shared guard, and no example copies a benchmark question.
"""

from __future__ import annotations

import json
import os
import re
import unittest
from datetime import date
from pathlib import Path

from chatbot.policies.kql_safety import inspect_query_safety
from chatbot.policies.plan_sanitization import sanitize_plan
from chatbot.policies.time_ranges import (
    EvaluationTimeContext, resolve_all_available_time_range, use_evaluation_time,
)
from chatbot.prompts import load_prompt, load_role_prompt

_WORKBOOK = Path(__file__).resolve().parents[1] / "Objective_Query_-_PHIA_v3.xlsx"


def _plan_examples() -> list[tuple[str, dict]]:
    system, _ = load_role_prompt("d1_sensor_data/plan_sensor_query.txt")
    text = system.replace("{{", "{").replace("}}", "}")
    return [(question, json.loads(body))
            for question, body in re.findall(r"\nQ: (.*?)\n(\{.*?\})\n(?=\n|Q:|If )", text, re.S)]


def _kql_examples() -> list[tuple[str, str]]:
    system, _ = load_role_prompt("d1_sensor_data/generate_sensor_query.txt")
    section = system.split("Two illustrations of the method", 1)[1]
    examples = []
    for block in re.split(r"\n\n(?=\S[^\n]*question: \")", section)[1:]:
        question = re.search(r'question: "(.*?)"', block, re.S).group(1)
        lines = block.splitlines()
        start = next(i for i, line in enumerate(lines) if re.match(r"(let |Sensor)", line))
        end = next((i for i in range(start, len(lines)) if not lines[i].strip()), len(lines))
        examples.append((" ".join(question.split()), "\n".join(lines[start:end])))
    return examples


class D1PromptExampleTests(unittest.TestCase):
    def setUp(self) -> None:
        os.environ.setdefault("KQL_GENERATE_METHOD", "llm")

    def test_prompts_keep_their_role_placeholders(self) -> None:
        for resource, feedback in (("plan_sensor_query.txt", "{validation_feedback}"),
                                   ("generate_sensor_query.txt", "{schema_observations}"),
                                   ("repair_sensor_query.txt", "{schema_observations}")):
            system, user = load_role_prompt("d1_sensor_data/" + resource)
            self.assertNotIn("{question}", system)
            self.assertIn("{question}", user)
            self.assertIn(feedback, user)
        load_prompt("d1_sensor_data/adx_schema_context.txt").format(user_id="u7")

    def test_planner_examples_are_executable_plans(self) -> None:
        examples = _plan_examples()
        self.assertGreaterEqual(len(examples), 3)
        clock = EvaluationTimeContext(
            summary_reference_date=date(2024, 3, 11), summary_min_date=date(2023, 11, 1),
            workout_reference_date=date(2024, 3, 11), workout_min_date=date(2023, 11, 1),
        )
        with use_evaluation_time(clock):
            for question, candidate in examples:
                with self.subTest(question=question):
                    if candidate["time_range"].get("mode") == "all_available":
                        candidate = {**candidate, "time_range": resolve_all_available_time_range(
                            candidate["time_range"], {**candidate, "timezone_iana": "UTC"})}
                    plan = sanitize_plan(candidate, question, {})
                    self.assertIsNot(plan.get("execution_ready"), False, plan.get("validation_errors"))
                    self.assertFalse(plan.get("validation_errors"))

    def test_query_examples_pass_the_shared_guard(self) -> None:
        examples = _kql_examples()
        self.assertGreaterEqual(len(examples), 2)
        for question, query in examples:
            with self.subTest(question=question):
                decision = inspect_query_safety(query, "u7")
                self.assertTrue(decision["is_safe"], decision["issues"])
                self.assertNotIn("```", query)

    def test_examples_are_not_benchmark_questions(self) -> None:
        questions = [q for q, _ in _plan_examples()] + [q for q, _ in _kql_examples()]
        for resource in ("plan_sensor_query.txt", "generate_sensor_query.txt",
                         "repair_sensor_query.txt", "adx_schema_context.txt"):
            self.assertNotIn("PHIA", load_prompt("d1_sensor_data/" + resource))
        if not _WORKBOOK.exists():
            self.skipTest("benchmark workbook not present")
        import openpyxl
        sheet = openpyxl.load_workbook(_WORKBOOK, read_only=True, data_only=True).worksheets[0]
        benchmark = {" ".join(str(row[0]).lower().split())
                     for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]}
        for question in questions:
            self.assertNotIn(" ".join(question.lower().split()), benchmark)


if __name__ == "__main__":
    unittest.main()
